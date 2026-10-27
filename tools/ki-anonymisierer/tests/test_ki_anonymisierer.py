import importlib.util
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import ki_anonymisierer as ka  # noqa: E402


def test_mapping_is_consistent_and_case_insensitive():
    m = ka.Mapping()
    assert m.token_for("PERSON", "Erika Mustermann") == "[PERSON_1]"
    assert m.token_for("PERSON", "erika  mustermann") == "[PERSON_1]"
    assert m.token_for("PERSON", "Max Mustermann") == "[PERSON_2]"
    assert m.token_for("ORT", "Berlin") == "[ORT_1]"


def test_mapping_roundtrip_on_disk(tmp_path):
    path = tmp_path / "z.json"
    m = ka.Mapping()
    m.token_for("PERSON", "Erika Mustermann")
    m.save(path)
    loaded = ka.Mapping.load(path)
    assert loaded.token_for("PERSON", "Erika Mustermann") == "[PERSON_1]"
    assert loaded.token_for("PERSON", "Neue Person") == "[PERSON_2]"


def test_restore_tolerates_ai_variations_and_reports_unknown():
    m = ka.Mapping()
    m.token_for("PERSON", "Erika Mustermann")
    m.token_for("ORT", "Berlin")
    text, unknown, unused = ka.restore("Frau [person_1] aus [ ORT_1 ] und [PERSON_7], [Hinweis_1]", m)
    assert text == "Frau Erika Mustermann aus Berlin und [PERSON_7], [Hinweis_1]"
    assert unknown == ["[PERSON_7]"]
    assert unused == []


class _R:
    def __init__(self, start, end, entity_type, score):
        self.start, self.end, self.entity_type, self.score = start, end, entity_type, score


def test_resolve_strips_titles_and_prefers_higher_score():
    text = "Dr. Thomas Weber, IBAN DE89 3704 0044 0532 0130 00."
    results = [
        _R(0, 16, "PERSON", 0.85),
        _R(18, 27, "LOCATION", 0.85),  # Fehlgriff des Sprachmodells, überlappt die IBAN
        _R(23, 50, "IBAN_CODE", 1.0),
    ]
    spans = ka._resolve(text, results)
    assert [(text[s.start : s.end], s.entity) for s in spans] == [
        ("Thomas Weber", "PERSON"),
        ("DE89 3704 0044 0532 0130 00", "IBAN_CODE"),
    ]


@pytest.mark.skipif(importlib.util.find_spec("de_core_news_md") is None, reason="spaCy-Modell fehlt")
def test_end_to_end_with_model():
    text = (
        "Frau Hildegard Schneider, wohnhaft Steinbachstraße 8a, 92361 Berngau, "
        "erreichbar unter pflege@example.de. Frau Schneider hat Pflegegrad 3."
    )
    mapping = ka.Mapping()
    result = ka.Anonymizer().anonymize(
        text, mapping, list(ka.LABELS), deny={"ORGANIZATION": []}, allow=["Pflegegrad"]
    )
    for secret in ("Schneider", "Steinbachstraße", "92361", "Berngau", "pflege@example.de"):
        assert secret not in result.text
    assert "Pflegegrad 3" in result.text
    restored, unknown, _ = ka.restore(result.text, mapping)
    assert restored == text
    assert unknown == []


# --------------------------------------------------------------------------- #
# Word: formaterhaltend
# --------------------------------------------------------------------------- #

from io import BytesIO  # noqa: E402

from docx import Document  # noqa: E402

import dokumente  # noqa: E402


def _docx_with_split_name() -> bytes:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Kopf: Erika Mustermann"
    p = doc.add_paragraph("Bewohnerin: Frau ")
    p.add_run("Eri").bold = True
    p.add_run("ka Muster").italic = True
    p.add_run("mann, Pflegegrad 3.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Erika Mustermann"
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_docx_replacement_keeps_formatting_across_runs():
    doc = dokumente.DocxText(_docx_with_split_name())
    text = doc.text
    replacements = []
    start = 0
    while (i := text.find("Erika Mustermann", start)) >= 0:
        replacements.append((i, i + len("Erika Mustermann"), "[PERSON_1]"))
        start = i + 1
    assert len(replacements) == 3  # Absatz, Tabelle, Kopfzeile
    doc.replace(replacements)

    result = Document(BytesIO(doc.save()))
    paragraph = result.paragraphs[0]
    assert paragraph.text == "Bewohnerin: Frau [PERSON_1], Pflegegrad 3."
    assert paragraph.runs[1].text == "[PERSON_1]" and paragraph.runs[1].bold
    assert result.tables[0].cell(0, 0).text == "[PERSON_1]"
    assert result.sections[0].header.paragraphs[0].text == "Kopf: [PERSON_1]"


def test_restore_document_docx_roundtrip():
    m = ka.Mapping()
    m.token_for("PERSON", "Erika Mustermann")
    doc = Document()
    doc.add_paragraph("Hallo ").add_run("[PERSON_1]").bold = True
    buffer = BytesIO()
    doc.save(buffer)
    result = ka.restore_document("antwort.docx", buffer.getvalue(), m)
    restored = Document(BytesIO(result.docx))
    assert restored.paragraphs[0].text == "Hallo Erika Mustermann"
    assert restored.paragraphs[0].runs[1].bold


def test_markdown_to_docx():
    data = dokumente.markdown_to_docx("# Titel\n\n- **fett** und *kursiv*\n1. eins\n\n| A | B |\n|---|---|\n| 1 | 2 |\n")
    doc = Document(BytesIO(data))
    assert [(p.style.name, p.text) for p in doc.paragraphs] == [
        ("Heading 1", "Titel"),
        ("List Bullet", "fett und kursiv"),
        ("List Number", "eins"),
    ]
    assert doc.paragraphs[1].runs[0].bold and doc.paragraphs[1].runs[2].italic
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [["A", "B"], ["1", "2"]]


# --------------------------------------------------------------------------- #
# Oberfläche (ohne Sprachmodell)
# --------------------------------------------------------------------------- #


@pytest.fixture
def server(tmp_path):
    import http.client
    import json
    import threading

    import oberflaeche

    app = oberflaeche.App(tmp_path / "arbeit", lambda: None)
    srv = oberflaeche.create_server(app, 0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    port = srv.server_address[1]

    def call(path, body=None, headers=None, host=None):
        conn = http.client.HTTPConnection("127.0.0.1", port)
        h = {"Host": host or f"127.0.0.1:{port}", "X-KI-Anonymisierer": "1", **(headers or {})}
        if body is None:
            conn.request("GET", path, headers=h)
        else:
            conn.request("POST", path, body=json.dumps(body), headers=h)
        response = conn.getresponse()
        raw = response.read()
        try:
            return response.status, json.loads(raw)
        except ValueError:
            return response.status, raw

    yield app, call
    srv.shutdown()


def test_server_rejects_foreign_host_and_missing_header(server):
    _app, call = server
    assert call("/api/status", host="evil.example:80")[0] == 403
    assert call("/api/vorgang/loeschen", {"vorgang": "x"}, headers={"X-KI-Anonymisierer": ""})[0] == 403


def test_server_restore_and_case_name_validation(server):
    app, call = server
    m = ka.Mapping()
    m.token_for("PERSON", "Erika Mustermann")
    m.save(app.cases / "Akte 1.zuordnung.json")

    status, data = call("/api/wiederherstellen", {"vorgang": "Akte 1", "text": "Frau [PERSON_1] und [PERSON_5]"})
    assert status == 200
    assert data["text"] == "Frau Erika Mustermann und [PERSON_5]"
    assert data["unbekannt"] == ["[PERSON_5]"]
    assert data["datei"]["name"].endswith(".docx")

    assert call("/api/wiederherstellen", {"vorgang": "../../etc/passwd", "text": "x"})[0] == 400
    status, data = call("/api/status")
    assert [c["name"] for c in data["vorgaenge"]] == ["Akte 1"]


def test_server_saves_terms(server):
    app, call = server
    call("/api/begriffe", {"immer": "ORGANISATION: Beispiel AG\nProjekt X", "nie": "Pflegegrad"})
    deny, allow = ka.load_terms(app.terms_path)
    assert deny == {"ORGANIZATION": ["Beispiel AG"], "BEGRIFF": ["Projekt X"]}
    assert allow == ["Pflegegrad"]
