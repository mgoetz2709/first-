"""Ein- und Ausgabe für den KI-Anonymisierer.

- Text aus .txt/.md/.docx/.pdf lesen
- Word-Dateien formaterhaltend bearbeiten: Ersetzungen erfolgen direkt in den
  Textstücken (w:t) der Datei, Schriftarten, Tabellen, Kopf-/Fußzeilen usw.
  bleiben erhalten — auch wenn ein Name über mehrere Formatierungen verteilt ist.
- KI-Antworten (Markdown) in eine einfach formatierte Word-Datei umwandeln
"""

from __future__ import annotations

import re
from io import BytesIO

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W_P = f"{{{W_NS}}}p"
W_T = f"{{{W_NS}}}t"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

Replacement = tuple[int, int, str]  # (start, ende, neuer Text) bezogen auf den Gesamttext


def apply_replacements(text: str, replacements: list[Replacement]) -> str:
    parts: list[str] = []
    pos = 0
    for start, end, new in sorted(replacements):
        parts.append(text[pos:start])
        parts.append(new)
        pos = end
    parts.append(text[pos:])
    return "".join(parts)


def read_text(name: str, data: bytes) -> str:
    """Liest den Text einer Datei anhand der Dateiendung."""
    suffix = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if suffix == "docx":
        return DocxText(data).text
    if suffix == "pdf":
        from pypdf import PdfReader

        return "\n\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252")


# --------------------------------------------------------------------------- #
# Word: formaterhaltende Ersetzung
# --------------------------------------------------------------------------- #


def _nearest_paragraph(element):
    parent = element.getparent()
    while parent is not None and parent.tag != W_P:
        parent = parent.getparent()
    return parent


def _replace_in_runs(texts, start: int, end: int, new: str) -> None:
    """Ersetzt [start, end) eines Absatzes, der über mehrere w:t-Elemente verteilt
    sein kann. Der neue Text landet im ersten betroffenen Stück und übernimmt
    dessen Formatierung; der Rest wird aus den folgenden Stücken entfernt."""
    pos = 0
    inserted = False
    for t in texts:
        content = t.text or ""
        a, b = pos, pos + len(content)
        pos = b
        if b <= start or a >= end:
            continue
        cut_from, cut_to = max(start, a) - a, min(end, b) - a
        t.text = content[:cut_from] + ("" if inserted else new) + content[cut_to:]
        t.set(XML_SPACE, "preserve")
        inserted = True


class DocxText:
    """Word-Dokument als durchgehender Text (ein Absatz pro Zeile), in dem
    Ersetzungen formaterhaltend in die Datei zurückgeschrieben werden."""

    def __init__(self, data: bytes):
        from docx import Document

        self.doc = Document(BytesIO(data))
        self.paragraphs: list[list] = []
        for root in self._roots():
            for p in root.iter(W_P):
                self.paragraphs.append([t for t in p.iter(W_T) if _nearest_paragraph(t) is p])

    def _roots(self):
        yield self.doc.element.body
        seen: set[int] = set()
        for section in self.doc.sections:
            for part in (
                section.header, section.first_page_header, section.even_page_header,
                section.footer, section.first_page_footer, section.even_page_footer,
            ):
                if part.is_linked_to_previous:
                    continue
                element = part._element
                if id(element) not in seen:
                    seen.add(id(element))
                    yield element

    @property
    def text(self) -> str:
        return "\n".join("".join(t.text or "" for t in texts) for texts in self.paragraphs)

    def replace(self, replacements: list[Replacement]) -> None:
        # Absatz-Grenzen im Gesamttext bestimmen
        bounds: list[tuple[int, int]] = []
        pos = 0
        for texts in self.paragraphs:
            length = sum(len(t.text or "") for t in texts)
            bounds.append((pos, pos + length))
            pos += length + 1  # "\n"

        per_paragraph: dict[int, list[Replacement]] = {}
        for start, end, new in replacements:
            first = True
            for i, (a, b) in enumerate(bounds):
                s, e = max(start, a), min(end, b)
                if s < e:  # Treffer, der über eine Absatzgrenze reicht, wird aufgeteilt
                    per_paragraph.setdefault(i, []).append((s - a, e - a, new if first else ""))
                    first = False

        for i, items in per_paragraph.items():
            for start, end, new in sorted(items, reverse=True):
                _replace_in_runs(self.paragraphs[i], start, end, new)

    def clear_metadata(self) -> None:
        """Entfernt Autor, Titel usw. aus den Dokumenteigenschaften."""
        props = self.doc.core_properties
        for name in ("author", "last_modified_by", "title", "subject", "keywords",
                     "comments", "category", "content_status", "identifier"):
            setattr(props, name, "")

    def save(self) -> bytes:
        buffer = BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Markdown (KI-Antwort) -> Word
# --------------------------------------------------------------------------- #

_INLINE_RE = re.compile(r"(\*\*[^*]+?\*\*|__[^_]+?__|\*[^*\s][^*]*?\*|`[^`]+`)")
_TABLE_SEPARATOR_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")


def _add_inline(paragraph, text: str) -> None:
    for piece in _INLINE_RE.split(text):
        if not piece:
            continue
        if (piece.startswith("**") and piece.endswith("**")) or (
            piece.startswith("__") and piece.endswith("__")
        ):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            paragraph.add_run(piece[1:-1]).font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def markdown_to_docx(text: str) -> bytes:
    """Wandelt einfaches Markdown (Überschriften, Listen, Tabellen, fett/kursiv)
    in eine Word-Datei um."""
    from docx import Document

    doc = Document()
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or re.fullmatch(r"[-*_]{3,}", stripped):
            i += 1
            continue
        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not _TABLE_SEPARATOR_RE.match(lines[i].strip()):
                    rows.append(_cells(lines[i]))
                i += 1
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell_paragraph = table.cell(r, c).paragraphs[0]
                    _add_inline(cell_paragraph, value)
                    if r == 0:
                        for run in cell_paragraph.runs:
                            run.bold = True
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        bullet = re.match(r"^[-*•]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if heading:
            paragraph = doc.add_heading(level=min(len(heading.group(1)), 4))
            _add_inline(paragraph, heading.group(2).strip("# "))
        elif bullet:
            _add_inline(doc.add_paragraph(style="List Bullet"), bullet.group(1))
        elif numbered:
            _add_inline(doc.add_paragraph(style="List Number"), numbered.group(1))
        else:
            _add_inline(doc.add_paragraph(), stripped)
        i += 1
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
